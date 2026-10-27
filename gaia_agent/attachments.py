"""Download and extract the attachment formats present in the GAIA subset."""

from __future__ import annotations

import csv
import mimetypes
import shutil
from pathlib import Path

import requests
from huggingface_hub import get_token, hf_hub_download

TEXT_SUFFIXES = {".txt", ".md", ".py", ".json", ".yaml", ".yml", ".xml"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
AUDIO_SUFFIXES = {".mp3", ".wav", ".m4a", ".ogg", ".webm"}
GAIA_DATASET_ID = "gaia-benchmark/GAIA"
GAIA_VALIDATION_DIR = "2023/validation"


def safe_name(file_name: str, fallback: str) -> str:
    name = Path(file_name or fallback).name
    if name in {"", ".", ".."}:
        raise ValueError("invalid attachment file name")
    return name


def download_attachment(
    api_url: str,
    task_id: str,
    file_name: str,
    task_dir: Path,
    timeout: int = 30,
) -> Path | None:
    if not file_name:
        return None
    task_dir.mkdir(parents=True, exist_ok=True)
    destination = task_dir / safe_name(file_name, task_id)
    if destination.exists() and destination.stat().st_size:
        return destination
    response = requests.get(f"{api_url}/files/{task_id}", timeout=timeout)
    if response.status_code == 404:
        token = get_token()
        if not token:
            return None
        cached_path = hf_hub_download(
            repo_id=GAIA_DATASET_ID,
            repo_type="dataset",
            filename=f"{GAIA_VALIDATION_DIR}/{destination.name}",
            token=token,
        )
        shutil.copyfile(cached_path, destination)
        return destination
    response.raise_for_status()
    destination.write_bytes(response.content)
    return destination


def list_task_files(task_dir: Path) -> list[str]:
    if not task_dir.exists():
        return []
    return sorted(path.name for path in task_dir.iterdir() if path.is_file())


def resolve_task_file(task_dir: Path, requested: str) -> Path:
    candidate = (task_dir / requested).resolve()
    root = task_dir.resolve()
    if root not in candidate.parents or not candidate.is_file():
        raise ValueError(f"file is unavailable: {requested!r}")
    return candidate


def _truncate(text: str, max_chars: int) -> str:
    if len(text) <= max_chars:
        return text
    return text[:max_chars] + f"\n...[truncated {len(text) - max_chars} characters]"


def read_attachment(path: Path, max_chars: int = 40_000) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _truncate(path.read_text(encoding="utf-8", errors="replace"), max_chars)
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = list(csv.reader(handle))
        return _truncate("\n".join("\t".join(row) for row in rows), max_chars)
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        text = "\n\n".join(
            f"--- page {index} ---\n{page.extract_text() or ''}"
            for index, page in enumerate(reader.pages, start=1)
        )
        return _truncate(text, max_chars)
    if suffix == ".docx":
        from docx import Document

        document = Document(path)
        return _truncate("\n".join(p.text for p in document.paragraphs), max_chars)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        chunks: list[str] = []
        for sheet in workbook.worksheets:
            chunks.append(f"--- sheet: {sheet.title} ---")
            rows = list(sheet.iter_rows(values_only=True))
            for row in rows:
                chunks.append("\t".join("" if value is None else str(value) for value in row))
                if sum(map(len, chunks)) > max_chars:
                    break
            if rows:
                headers = rows[0]
                totals: list[str] = []
                for column, header in enumerate(headers):
                    values = [
                        row[column]
                        for row in rows[1:]
                        if column < len(row)
                        and isinstance(row[column], (int, float))
                        and not isinstance(row[column], bool)
                    ]
                    if values:
                        total = sum(values)
                        rendered = str(int(total)) if float(total).is_integer() else str(total)
                        totals.append(f"{header or f'column {column + 1}'}\t{rendered}")
                if totals:
                    chunks.append("--- numeric column totals (all data rows) ---")
                    chunks.extend(totals)
        workbook.close()
        return _truncate("\n".join(chunks), max_chars)
    mime = mimetypes.guess_type(path.name)[0] or "unknown"
    if suffix in IMAGE_SUFFIXES:
        return f"Image file {path.name} ({mime}, {path.stat().st_size} bytes) is attached to the user message."
    if suffix in AUDIO_SUFFIXES:
        return f"Audio file {path.name} ({mime}, {path.stat().st_size} bytes); use the transcript in the user message."
    raise ValueError(f"unsupported attachment type: {suffix or '<none>'}")


def data_url(path: Path) -> str:
    import base64

    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    payload = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{payload}"
